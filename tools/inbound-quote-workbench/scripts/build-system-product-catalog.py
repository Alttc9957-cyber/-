#!/usr/bin/env python3
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


INPUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/Users/alic/Downloads/产品库汇总.xlsx")
OUTPUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parents[1] / "data/products/youyixing-product-catalog.json"
PENDING_SUPPLIER = "待绑定供应商"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def price(value):
    value = clean(value)
    if value == "":
        return ""
    if isinstance(value, (int, float)):
        return value if value == value else ""
    text = str(value).strip().replace("¥", "").replace("￥", "").replace(",", "").replace(" ", "")
    return float(text) if re.fullmatch(r"-?\d+(?:\.\d+)?", text) else ""


def first_present(*values):
    for value in values:
        if value != "" and value is not None:
            return value
    return ""


def norm_city(value):
    value = str(clean(value)).replace("市", "").strip()
    value = re.sub(r".*[+＋、，,]", "", value)
    aliases = {
        "beijing": "北京",
        "chongqing": "重庆",
        "chengdu": "成都",
        "kunming": "昆明",
        "shanghai": "上海",
        "xian": "西安",
        "xi an": "西安",
        "hangzhou": "杭州",
    }
    return aliases.get(value.lower(), value)


def norm_vehicle_type(value):
    text = str(clean(value))
    if re.search(r"接送机|接机|送机|机场|大兴|首都机场|airport", text, re.I):
        return "接送机"
    if re.search(r"接送站|接站|送站|火车站|高铁站|station|train", text, re.I):
        return "接送站"
    if re.search(r"包车|一日游|两日游|三日游|用车|市区|市内|本地游|8小时|9小时|武隆", text):
        return "包车"
    return text or "包车"


def norm_model(value):
    text = str(clean(value)).replace("车", "").replace(" ", "")
    if re.search(r"14|15|16|17", text):
        return "14座~17座"
    if re.search(r"36|38|33", text):
        return "36~38座"
    if re.search(r"51|55|50", text):
        return "51~55座"
    if "22" in text:
        return "22座"
    if "9" in text:
        return "9座"
    if re.search(r"7|8", text):
        return "7座"
    if "5" in text:
        return "5座"
    return text


def norm_language(value):
    text = str(clean(value))
    if not text:
        return "英语"
    if re.search(r"英|english", text, re.I):
        return "英语"
    if re.search(r"西|spanish|español", text, re.I):
        return "西语"
    if re.search(r"法|french", text, re.I):
        return "法语"
    if re.search(r"德|german", text, re.I):
        return "德语"
    if re.search(r"意|italian", text, re.I):
        return "意语"
    if re.search(r"日|japanese", text, re.I):
        return "日语"
    if re.search(r"韩|korean", text, re.I):
        return "韩语"
    if re.search(r"俄|russian", text, re.I):
        return "俄语"
    return text


def norm_star(value):
    text = str(clean(value))
    if "豪华" in text:
        return "五星级/豪华型"
    if re.search(r"5|五", text):
        return "五星级"
    if re.search(r"4|四|舒适", text):
        return "四星级"
    if re.search(r"3|三", text):
        return "三星级"
    return text or "待补星级"


def norm_room(value):
    text = str(clean(value))
    if re.search(r"双|标|大床或标间|大床或标", text):
        return "双床房"
    if "大床" in text:
        return "大床房"
    if "三" in text:
        return "三人间"
    if "套" in text:
        return "套房"
    return text or "房型待定"


def row_values(ws):
    for row in ws.iter_rows(values_only=True):
        yield [clean(cell) for cell in row]


def raw(labels, row):
    return {label: clean(row[i]) if i < len(row) else "" for i, label in enumerate(labels) if label}


def merged_value(ws, row, col):
    value = ws.cell(row, col).value
    if clean(value) != "":
        return clean(value)
    for cell_range in ws.merged_cells.ranges:
        if cell_range.min_row <= row <= cell_range.max_row and cell_range.min_col <= col <= cell_range.max_col:
            return clean(ws.cell(cell_range.min_row, cell_range.min_col).value)
    return ""


def compact_header(value):
    return re.sub(r"\s+", " ", str(clean(value))).strip()


def header_labels(ws, header_rows):
    labels = []
    for col in range(1, ws.max_column + 1):
        parts = []
        for row in header_rows:
            value = compact_header(merged_value(ws, row, col))
            if value and value not in parts:
                parts.append(value)
        labels.append(" / ".join(parts) or f"{get_column_letter(col)}列")
    return labels


def raw_fields(ws, row_number, labels, header_rows):
    row = [clean(ws.cell(row_number, col).value) for col in range(1, ws.max_column + 1)]
    fields = raw(labels, row)
    for col, label in enumerate(header_labels(ws, header_rows), start=1):
        fields[f"{get_column_letter(col)}列 / {label}"] = row[col - 1] if col - 1 < len(row) else ""
    return fields


def item_id(category, sheet, row_number, suffix=""):
    safe_sheet = re.sub(r"\W+", "-", sheet).strip("-")
    return f"SYS-{category}-{safe_sheet}-{row_number}{suffix}"


def status_for_cost(cost):
    return "可报价" if cost != "" else "待补成本"


def parse_routes(ws, report):
    labels = [
        "产品编号", "产品类型", "产品城市", "线路天数", "产品名", "产品详情链接", "产品展示信息", "报价包含",
        "C端/国外B端2~3人", "C端/国外B端4~5人", "C端/国外B端6~8人", "C端/国外B端9~13人", "C端/国外B端14~18人",
        "国内B端2~3人", "国内B端4~5人", "国内B端6~8人", "国内B端9~13人", "国内B端14~18人",
        "", "成本2~3人", "成本4~5人", "成本6~8人", "成本9~13人", "成本14~18人", "成本详情链接"
    ]
    tier_labels = ["2~3人", "4~5人", "6~8人", "9~13人", "14~18人"]
    out = []
    for idx, row in enumerate(list(row_values(ws))[2:], start=3):
        if not any(str(x).strip() for x in row):
            report["skippedEmptyRows"] += 1
            continue
        if not clean(row[4] if len(row) > 4 else ""):
            report["skippedEmptyRows"] += 1
            continue
        cost_tiers = {label: price(row[19 + i] if len(row) > 19 + i else "") for i, label in enumerate(tier_labels)}
        item = {
            "id": item_id("线路产品", ws.title, idx),
            "category": "线路产品",
            "code": clean(row[0] if len(row) > 0 else ""),
            "productType": clean(row[1] if len(row) > 1 else ""),
            "city": norm_city(row[2] if len(row) > 2 else ""),
            "days": price(row[3] if len(row) > 3 else ""),
            "name": clean(row[4] if len(row) > 4 else ""),
            "detailLink": clean(row[5] if len(row) > 5 else ""),
            "displayLink": clean(row[6] if len(row) > 6 else ""),
            "includes": [x.strip() for x in re.split(r"[、,，/]+", str(clean(row[7] if len(row) > 7 else ""))) if x.strip()],
            "priceTiers": {
                "foreignSale": {label: price(row[8 + i] if len(row) > 8 + i else "") for i, label in enumerate(tier_labels)},
                "domesticSale": {label: price(row[13 + i] if len(row) > 13 + i else "") for i, label in enumerate(tier_labels)},
                "cost": cost_tiers,
            },
            "costTiers": cost_tiers,
            "costDetailLink": clean(row[24] if len(row) > 24 else ""),
            "costPrice": first_present(*cost_tiers.values()),
            "supplierName": PENDING_SUPPLIER,
            "status": "可报价" if any(v != "" for v in cost_tiers.values()) else "待补成本",
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1, 2]),
            "importWarnings": [],
        }
        out.append(item)
    return out


def parse_vehicles(ws, report):
    labels = ["城市", "订单类型", "行程", "卖价5座", "卖价7座", "卖价9座", "卖价14座~17座", "卖价22座", "卖价36~38座", "卖价51~55座", "司机兼导游卖价5座", "司机兼导游卖价7座", "空列", "成本5座", "成本7座", "成本9座", "成本14座~17座", "成本22座", "成本36~38座", "成本51~55座", "司机兼导游成本5座", "司机兼导游成本7座"]
    models = ["5座", "7座", "9座", "14座~17座", "22座", "36~38座", "51~55座"]
    out, city = [], ""
    for idx, row in enumerate(list(row_values(ws))[6:], start=7):
        if not any(str(x).strip() for x in row):
            report["skippedEmptyRows"] += 1
            continue
        if len(row) > 0 and row[0]:
            city = norm_city(row[0])
        service_type = norm_vehicle_type(first_present(row[1] if len(row) > 1 else "", row[2] if len(row) > 2 else ""))
        route = clean(first_present(row[2] if len(row) > 2 else "", row[1] if len(row) > 1 else "", service_type))
        if not city or not service_type:
            report["abnormalRows"] += 1
            report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或服务类型"})
            continue
        for mi, model in enumerate(models):
            sale = price(row[3 + mi] if len(row) > 3 + mi else "")
            cost = price(row[13 + mi] if len(row) > 13 + mi else "")
            if sale == "" and cost == "":
                continue
            model = norm_model(model)
            item = {
                "id": item_id("用车", ws.title, idx, f"-{model}"),
                "category": "用车",
                "city": city,
                "name": f"{city}{service_type} {route} {model}",
                "serviceType": service_type,
                "vehicleType": service_type,
                "route": route,
                "model": model,
                "seatCount": int(re.search(r"\d+", model).group()) if re.search(r"\d+", model) else "",
                "costPrice": cost,
                "salePrice": sale,
                "dayCost": cost if service_type == "包车" else "",
                "airportTransferCost": cost if service_type == "接送机" else "",
                "stationTransferCost": cost if service_type == "接送站" else "",
                "airportRoute": route if service_type == "接送机" else "",
                "stationRoute": route if service_type == "接送站" else "",
                "pricingUnit": "per_trip" if service_type in ("接送机", "接送站") else "per_day",
                "supplierName": PENDING_SUPPLIER,
                "status": status_for_cost(cost),
                "source": "Excel系统产品库",
                "sourceSheet": ws.title,
                "sourceRow": idx,
                "rawFields": raw_fields(ws, idx, labels, [1, 2, 3, 4, 5, 6]),
                "importWarnings": [],
            }
            out.append(item)
    return out


def parse_guides(ws, report):
    labels = ["城市", "语种", "淡季卖价", "淡季超时费", "淡季住宿费", "旺季卖价", "旺季超时费", "旺季住宿费", "导游是否需要门票", "空列", "淡季成本", "淡季成本超时费", "淡季成本住宿费", "旺季成本", "旺季成本超时费", "旺季成本住宿费"]
    out, city = [], ""
    for idx, row in enumerate(list(row_values(ws))[2:], start=3):
        if not any(str(x).strip() for x in row):
            report["skippedEmptyRows"] += 1
            continue
        if len(row) > 0 and row[0]:
            city = norm_city(row[0])
        language = norm_language(row[1] if len(row) > 1 else "")
        if not city or not language:
            report["abnormalRows"] += 1
            report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或语种"})
            continue
        low_cost = price(row[10] if len(row) > 10 else "")
        high_cost = price(row[13] if len(row) > 13 else "")
        item = {
            "id": item_id("导游", ws.title, idx),
            "category": "导游",
            "city": city,
            "name": f"{city}{language}导游",
            "language": language,
            "lowSeasonSale": price(row[2] if len(row) > 2 else ""),
            "lowSeasonOvertimeFee": price(row[3] if len(row) > 3 else ""),
            "lowSeasonAccommodation": price(row[4] if len(row) > 4 else ""),
            "highSeasonSale": price(row[5] if len(row) > 5 else ""),
            "highSeasonOvertimeFee": price(row[6] if len(row) > 6 else ""),
            "highSeasonAccommodation": price(row[7] if len(row) > 7 else ""),
            "needsTicket": clean(row[8] if len(row) > 8 else ""),
            "lowSeasonCost": low_cost,
            "lowSeasonCostOvertime": price(row[11] if len(row) > 11 else ""),
            "lowSeasonCostAccommodation": price(row[12] if len(row) > 12 else ""),
            "highSeasonCost": high_cost,
            "highSeasonCostOvertime": price(row[14] if len(row) > 14 else ""),
            "highSeasonCostAccommodation": price(row[15] if len(row) > 15 else ""),
            "fullDayCost": first_present(low_cost, high_cost),
            "costPrice": first_present(low_cost, high_cost),
            "salePrice": first_present(price(row[2] if len(row) > 2 else ""), price(row[5] if len(row) > 5 else "")),
            "supplierName": PENDING_SUPPLIER,
            "status": status_for_cost(first_present(low_cost, high_cost)),
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1, 2]),
            "importWarnings": [],
        }
        out.append(item)
    return out


def parse_experiences(ws, report):
    labels = ["城市", "体验名称", "票种", "体验时间", "体验介绍信息", "成人卖价", "儿童卖价", "官方成人价", "官方儿童价", "成人成本", "儿童成本", "空列", "备注"]
    out, city, exp = [], "", ""
    for idx, row in enumerate(list(row_values(ws))[1:], start=2):
        if len(row) > 0 and row[0]:
            city = norm_city(row[0])
        if len(row) > 1 and row[1]:
            exp = clean(row[1])
        ticket_type = clean(row[2] if len(row) > 2 else "") or "体验项目"
        if not city or not exp:
            if any(str(x).strip() for x in row):
                report["abnormalRows"] += 1
                report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或体验名称"})
            continue
        adult_cost = price(row[9] if len(row) > 9 else "")
        item = {
            "id": item_id("特色体验", ws.title, idx),
            "category": "特色体验",
            "city": city,
            "name": exp if ticket_type == "体验项目" else f"{exp}{ticket_type}",
            "experienceName": exp,
            "ticketType": ticket_type,
            "duration": clean(row[3] if len(row) > 3 else ""),
            "introLink": clean(row[4] if len(row) > 4 else ""),
            "adultSale": price(row[5] if len(row) > 5 else ""),
            "childSale": price(row[6] if len(row) > 6 else ""),
            "officialPrice": price(row[7] if len(row) > 7 else ""),
            "officialChildPrice": price(row[8] if len(row) > 8 else ""),
            "adultCost": adult_cost,
            "childCost": price(row[10] if len(row) > 10 else ""),
            "costPrice": adult_cost,
            "salePrice": price(row[5] if len(row) > 5 else ""),
            "remark": clean(row[12] if len(row) > 12 else ""),
            "supplierName": PENDING_SUPPLIER,
            "status": status_for_cost(adult_cost),
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1]),
            "importWarnings": [],
        }
        out.append(item)
    return out


def parse_tickets(ws, report):
    labels = ["城市", "景点名称", "类型", "票种", "淡季成人", "淡季儿童", "旺季成人", "旺季儿童", "旅行社成人", "旅行社儿童", "免费政策", "备注", "保票政策", "保票联系人", "保票电话"]
    out, city, scenic = [], "", ""
    for idx, row in enumerate(list(row_values(ws))[2:], start=3):
        if len(row) > 0 and row[0]:
            city = norm_city(row[0])
        if len(row) > 1 and row[1]:
            scenic = clean(row[1])
        ticket_type = clean(row[3] if len(row) > 3 else "") or clean(row[2] if len(row) > 2 else "") or "景区门票"
        if not city or not scenic:
            if any(str(x).strip() for x in row):
                report["abnormalRows"] += 1
                report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或景点名称"})
            continue
        agency_adult = price(row[8] if len(row) > 8 else "")
        peak_adult = price(row[6] if len(row) > 6 else "")
        off_adult = price(row[4] if len(row) > 4 else "")
        cost = first_present(agency_adult, peak_adult, off_adult)
        item = {
            "id": item_id("景点门票", ws.title, idx),
            "category": "景点门票",
            "city": city,
            "name": scenic,
            "scenicName": scenic,
            "attractionLevel": clean(row[2] if len(row) > 2 else ""),
            "ticketType": ticket_type,
            "offAdult": off_adult,
            "offChild": price(row[5] if len(row) > 5 else ""),
            "offDiscount": price(row[5] if len(row) > 5 else ""),
            "peakAdult": peak_adult,
            "peakChild": price(row[7] if len(row) > 7 else ""),
            "peakDiscount": price(row[7] if len(row) > 7 else ""),
            "agencyAdult": agency_adult,
            "agencyChild": price(row[9] if len(row) > 9 else ""),
            "agencyDiscount": price(row[9] if len(row) > 9 else ""),
            "freePolicy": clean(row[10] if len(row) > 10 else ""),
            "remark": clean(row[11] if len(row) > 11 else ""),
            "guaranteePolicy": clean(row[12] if len(row) > 12 else ""),
            "guaranteeContactName": clean(row[13] if len(row) > 13 else ""),
            "guaranteeContactPhone": clean(row[14] if len(row) > 14 else ""),
            "costPrice": cost,
            "salePrice": first_present(peak_adult, off_adult, agency_adult),
            "supplierName": PENDING_SUPPLIER,
            "status": status_for_cost(cost),
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1, 2]),
            "importWarnings": [],
        }
        out.append(item)
    return out


def parse_meals(ws, report):
    labels = ["城市", "菜品", "是否清真", "饭店", "人均最低卖价", "电话", "地段", "空列", "建议卖价", "人均最低成本", "建议卖价2", "加价或备注"]
    out, city, cuisine, halal = [], "", "", ""
    for idx, row in enumerate(list(row_values(ws))[1:], start=2):
        if len(row) > 0 and row[0]:
            city = norm_city(row[0])
        if len(row) > 1 and row[1]:
            cuisine = clean(row[1])
        if len(row) > 2 and row[2]:
            halal = clean(row[2])
        restaurant = clean(row[3] if len(row) > 3 else "")
        if not city or not restaurant:
            if any(str(x).strip() for x in row):
                report["abnormalRows"] += 1
                report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或饭店"})
            continue
        cost = price(row[9] if len(row) > 9 else "")
        item = {
            "id": item_id("餐厅", ws.title, idx),
            "category": "餐厅",
            "city": city,
            "name": restaurant,
            "cuisine": cuisine,
            "halal": halal,
            "restaurant": restaurant,
            "minSale": price(row[4] if len(row) > 4 else ""),
            "phone": clean(row[5] if len(row) > 5 else ""),
            "area": clean(row[6] if len(row) > 6 else ""),
            "suggestedSale": price(row[8] if len(row) > 8 else ""),
            "suggestedSaleAlt": price(row[10] if len(row) > 10 else ""),
            "markup": price(row[11] if len(row) > 11 else ""),
            "minCost": cost,
            "costPrice": cost,
            "salePrice": first_present(price(row[8] if len(row) > 8 else ""), price(row[4] if len(row) > 4 else "")),
            "remark": clean(row[11] if len(row) > 11 else ""),
            "supplierName": PENDING_SUPPLIER,
            "status": status_for_cost(cost),
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1]),
            "importWarnings": [],
        }
        out.append(item)
    return out


def parse_hotels(ws, report):
    labels = ["城市", "酒店名称", "星级", "位置优势", "携程链接", "是否有协议", "协议固定", "房型", "卖价", "早餐", "备注", "空列1", "空列2", "成本"]
    out = []
    for idx, row in enumerate(list(row_values(ws))[1:], start=2):
        city = norm_city(row[0] if len(row) > 0 else "")
        hotel = clean(row[1] if len(row) > 1 else "")
        if not city or not hotel:
            if any(str(x).strip() for x in row):
                report["abnormalRows"] += 1
                report["issues"].append({"sheet": ws.title, "rowNumber": idx, "message": "缺城市或酒店名称"})
            continue
        cost = price(row[13] if len(row) > 13 else "")
        item = {
            "id": item_id("酒店", ws.title, idx),
            "category": "酒店",
            "city": city,
            "name": hotel,
            "hotelName": hotel,
            "star": norm_star(row[2] if len(row) > 2 else ""),
            "locationAdvantage": clean(row[3] if len(row) > 3 else ""),
            "ctripLink": clean(row[4] if len(row) > 4 else ""),
            "hasAgreement": clean(row[5] if len(row) > 5 else ""),
            "agreementFixed": clean(row[6] if len(row) > 6 else ""),
            "roomType": norm_room(row[7] if len(row) > 7 else ""),
            "salePrice": price(row[8] if len(row) > 8 else ""),
            "breakfast": clean(row[9] if len(row) > 9 else ""),
            "remark": clean(row[10] if len(row) > 10 else ""),
            "costPrice": cost,
            "agreementCost": cost,
            "supplierName": PENDING_SUPPLIER,
            "status": status_for_cost(cost),
            "source": "Excel系统产品库",
            "sourceSheet": ws.title,
            "sourceRow": idx,
            "rawFields": raw_fields(ws, idx, labels, [1]),
            "importWarnings": [],
        }
        out.append(item)
    return out


PARSERS = {
    "线路报价": ("routes", parse_routes),
    "仅包车报价": ("vehicles", parse_vehicles),
    "导游报价": ("guides", parse_guides),
    "特色体验价": ("experiences", parse_experiences),
    "门票报价": ("tickets", parse_tickets),
    "餐": ("meals", parse_meals),
    "酒店": ("hotels", parse_hotels),
}


def main():
    wb = openpyxl.load_workbook(INPUT, data_only=True)
    catalog = {key: [] for key in ["routes", "vehicles", "experiences", "tickets", "guides", "hotels", "meals", "transports", "others"]}
    report = {"sourceFile": str(INPUT), "sheets": {}, "skippedEmptyRows": 0, "abnormalRows": 0, "issues": []}
    for sheet_name, (key, parser) in PARSERS.items():
        ws = wb[sheet_name]
        before_skip = report["skippedEmptyRows"]
        before_abnormal = report["abnormalRows"]
        items = parser(ws, report)
        catalog[key].extend(items)
        report["sheets"][sheet_name] = {
            "sheetName": sheet_name,
            "category": {"routes": "线路产品", "vehicles": "用车", "guides": "导游", "experiences": "特色体验", "tickets": "景点门票", "meals": "餐厅", "hotels": "酒店"}[key],
            "recognized": True,
            "imported": len(items),
            "skippedEmptyRows": report["skippedEmptyRows"] - before_skip,
            "abnormalRows": report["abnormalRows"] - before_abnormal,
        }
    payload = {
        "meta": {
            "sourceFile": str(INPUT),
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "mode": "system_excel_overwrite",
            "version": "product-catalog-system-20260703",
        },
        "productCatalog": catalog,
        "report": report,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(OUTPUT),
        "counts": {key: len(value) for key, value in catalog.items()},
        "sheets": report["sheets"],
        "skippedEmptyRows": report["skippedEmptyRows"],
        "abnormalRows": report["abnormalRows"],
        "issues": report["issues"][:10],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
